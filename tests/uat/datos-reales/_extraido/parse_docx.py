"""Extrae preguntas de los .docx de origen a JSON estructurado.

Uso: python3 parse_docx.py
Lee los .docx de tests/uat/datos-reales/ y escribe
tests/uat/datos-reales/preguntas.json — revisar a mano antes de sembrar.

Los 3 cuestionarios de Gestión de Proyectos no marcan la opción correcta con un patrón
sintáctico simple (negrita+rojo sobre texto libre, con párrafos de "escenario" intercalados
entre la pregunta y las opciones) — se relevaron a mano los índices de párrafo no vacíos de
cada archivo (con un script auxiliar) y se codifican aquí como datos explícitos, no heurística.
Diseño/Parte 1/Parte 2 sí tienen un marcador confiable por pregunta y se parsean en forma
genérica.
"""

from __future__ import annotations

import json
import re
from pathlib import Path

import docx

BASE = Path(__file__).resolve().parent.parent
LETRA_RE = re.compile(r"^[\s]*[-•]?\s*([a-dA-D])[\)\.]\s*(.*)$", re.DOTALL)


def dificultad_de(texto: str) -> str:
    if len(texto) > 220:
        return "alto"
    if len(texto) < 60:
        return "bajo"
    return "medio"


def limpiar_opcion(texto: str) -> str:
    m = LETRA_RE.match(texto)
    return m.group(2).strip() if m else texto.strip()


def nonempty_paragraphs(path: Path) -> list[str]:
    d = docx.Document(path)
    return [p.text.strip() for p in d.paragraphs if p.text.strip()]


def armar_om(materia, unidad, tema, texto, opciones, correcta_idx, dificultad=None):
    return {
        "materia": materia,
        "unidad_tematica": unidad,
        "tema": tema,
        "tipo": "opcion_multiple",
        "texto": texto.strip(),
        "opciones": [limpiar_opcion(o) for o in opciones],
        "indice_correcta": correcta_idx,
        "dificultad": dificultad or dificultad_de(texto),
        "importancia": "medio",
    }


def armar_vf(materia, unidad, tema, texto, respuesta_correcta: bool, dificultad=None):
    return {
        "materia": materia,
        "unidad_tematica": unidad,
        "tema": tema,
        "tipo": "verdadero_falso",
        "texto": texto.strip(),
        "respuesta_correcta": respuesta_correcta,
        "dificultad": dificultad or dificultad_de(texto),
        "importancia": "medio",
    }


# ---------------------------------------------------------------------------
# Gestión de Proyectos — datos explícitos por índice de párrafo (ver docstring)
# ---------------------------------------------------------------------------


def parse_cuestionario_2(path: Path) -> list[dict]:
    p = nonempty_paragraphs(path)
    materia = "Gestión de Proyectos"
    out = []
    u1, t1 = "Ciclo de vida del proyecto", "Fases del proyecto"
    out.append(armar_om(materia, u1, t1, p[0], p[1:5], 1))
    out.append(armar_om(materia, u1, t1, p[5], p[6:10], 0))
    u2, t2 = "Ciclo de vida del proyecto", "Modelos de ciclo de vida (predictivo/adaptativo)"
    out.append(armar_om(materia, u2, t2, p[10], p[11:14], 1))
    out.append(armar_om(materia, u2, t2, p[14], p[15:18], 2))
    out.append(armar_om(materia, u2, t2, p[18], p[19:22], 0))
    u3, t3 = "Metodologías de desarrollo de software", "Selección de metodología según el contexto"
    out.append(armar_om(materia, u3, t3, f"{p[22]} {p[23]}", p[24:28], 2))
    out.append(armar_om(materia, u3, t3, p[28], p[29:33], 0))
    out.append(armar_om(materia, u3, t3, f"{p[33]} {p[34]}", p[35:39], 1))
    out.append(armar_om(materia, u3, t3, f"{p[39]} {p[40]}", p[41:45], 2))
    u4, t4 = "Liderazgo y equipos", "Liderazgo servicial y motivación"
    out.append(armar_om(materia, u4, t4, p[45], p[46:50], 1))
    out.append(armar_om(materia, u4, t4, p[51], p[52:56], 1))
    u5, t5 = "Liderazgo y equipos", "Dirección de proyectos"
    out.append(armar_om(materia, u5, t5, p[56], p[57:60], 1))
    return out


def parse_cuestionario_3(path: Path) -> list[dict]:
    p = nonempty_paragraphs(path)
    materia = "Gestión de Proyectos"
    out = []
    u1, t1 = "Planificación y cronograma", "EDT y cronograma"
    out.append(armar_om(materia, u1, t1, p[0], p[1:4], 0))
    out.append(armar_om(materia, u1, t1, p[4], p[5:9], 3))
    out.append(armar_om(materia, u1, t1, p[9], p[10:14], 0))
    out.append(armar_om(materia, u1, t1, p[14], p[15:19], 2))
    u2, t2 = "Planificación y cronograma", "Ruta crítica"
    out.append(armar_vf(materia, u2, t2, p[19], True))
    out.append(armar_om(materia, u2, t2, p[22], p[23:26], 0))
    u3, t3 = "Estimación ágil", "Técnicas de estimación"
    out.append(armar_om(materia, u3, t3, p[26], p[27:31], 1))
    out.append(armar_om(materia, u3, t3, p[31], p[32:36], 1))
    out.append(armar_om(materia, u3, t3, p[36], p[37:40], 2))
    out.append(armar_om(materia, u3, t3, p[40], p[41:44], 2))
    out.append(armar_om(materia, u3, t3, p[44], p[45:48], 1))
    out.append(armar_vf(materia, u1, t1, p[48], False))
    u4, t4 = "Requisitos", "Elicitación de requisitos"
    out.append(armar_vf(materia, u4, t4, p[51], True))
    out.append(armar_vf(materia, u4, t4, p[54], True))
    out.append(armar_vf(materia, u4, t4, p[57], False))
    return out


def parse_cuestionario_4(path: Path) -> list[dict]:
    p = nonempty_paragraphs(path)
    materia = "Gestión de Proyectos"
    out = []
    u1, t1 = "Estructuras organizacionales", "Tipos de estructura organizacional"
    out.append(armar_om(materia, u1, t1, p[0], p[1:4], 0))
    out.append(armar_om(materia, u1, t1, p[4], p[5:8], 2))
    u2, t2 = "Roles y responsabilidades", "Matriz RACI"
    out.append(armar_om(materia, u2, t2, p[8], p[9:13], 3))
    out.append(armar_om(materia, u2, t2, p[13], p[14:18], 1))
    u3, t3 = "Roles y responsabilidades", "Roles de Scrum"
    out.append(armar_om(materia, u3, t3, p[18], p[19:23], 2))
    out.append(armar_om(materia, u3, t3, p[23], p[24:28], 1))
    out.append(armar_om(materia, u3, t3, p[28], p[29:33], 3))
    out.append(armar_om(materia, u3, t3, p[33], p[34:38], 1))
    u4, t4 = "Interesados y comunicación", "Gestión de interesados y comunicaciones"
    out.append(armar_om(materia, u4, t4, p[38], p[39:43], 1))
    out.append(armar_om(materia, u4, t4, p[43], p[44:48], 1))
    out.append(armar_om(materia, u4, t4, p[48], p[49:53], 0))
    return out


# ---------------------------------------------------------------------------
# Ingeniería de Software — parsers genéricos (marcador confiable por pregunta)
# ---------------------------------------------------------------------------


def parse_bold_marks(path: Path, materia: str, unidad: str, tema: str) -> list[dict]:
    """Parte 1 / Parte 2: negrita simple (sin color) marca la opción correcta."""
    d = docx.Document(path)
    preguntas: list[dict] = []
    actual_texto = None
    actual_opciones: list[str] = []
    actual_correcta = None

    def flush():
        nonlocal actual_texto, actual_opciones, actual_correcta
        if actual_texto and actual_opciones:
            correcta = actual_correcta if actual_correcta is not None else 0
            preguntas.append(armar_om(materia, unidad, tema, actual_texto, actual_opciones, correcta))
        actual_texto, actual_opciones, actual_correcta = None, [], None

    for p in d.paragraphs:
        texto = p.text.strip()
        if not texto:
            continue
        if not LETRA_RE.match(texto):
            flush()
            actual_texto = re.sub(r"^\d+\.\s*", "", texto)
            continue
        actual_opciones.append(texto)
        if any(r.bold for r in p.runs):
            actual_correcta = len(actual_opciones) - 1
    flush()
    return preguntas


def parse_diseno(path: Path, materia: str) -> list[dict]:
    """Preguntas - Diseño.docx: un párrafo por pregunta completa, con '✅ Respuesta correcta: X)'."""
    d = docx.Document(path)
    secciones = [
        (0, "Principios de diseño", "Principios generales (KISS, DRY, YAGNI...)"),
        (7, "Deuda técnica", "Deuda técnica y leyes de Lehman"),
        (14, "Cohesión y acoplamiento", "Cohesión"),
        (19, "Cohesión y acoplamiento", "Acoplamiento"),
        (24, "Ocultamiento de información", "Ocultamiento de información"),
        (29, "Principios SOLID", "SOLID"),
    ]
    preguntas: list[dict] = []
    for p in d.paragraphs:
        texto = p.text.strip()
        if not texto or "✅" not in texto:
            continue
        cuerpo, respuesta = texto.split("✅ Respuesta correcta:")
        m_resp = re.match(r"^\s*([a-dA-D])\)", respuesta.strip())
        letra_correcta = m_resp.group(1).lower() if m_resp else None

        lineas = [l.strip() for l in cuerpo.strip().split("\n") if l.strip()]
        pregunta_texto = re.sub(r"^\d+\.\s*", "", lineas[0])
        opciones = lineas[1:]
        correcta_idx = 0
        for i, o in enumerate(opciones):
            m = LETRA_RE.match(o)
            if m and letra_correcta and m.group(1).lower() == letra_correcta:
                correcta_idx = i
                break

        idx = len(preguntas)
        unidad, tema = "Sin clasificar", "Sin clasificar"
        for umbral, u, t in secciones:
            if idx >= umbral:
                unidad, tema = u, t
        preguntas.append(armar_om(materia, unidad, tema, pregunta_texto, opciones, correcta_idx))
    return preguntas


def main() -> None:
    todas: list[dict] = []

    todas += parse_cuestionario_2(BASE / "Gestión de Proyectos -  Cuestionario 2.docx")
    todas += parse_cuestionario_3(BASE / "Gestión de Proyectos -  Cuestionario 3.docx")
    todas += parse_cuestionario_4(BASE / "Gestión de Proyectos -  Cuestionario 4.docx")

    todas += parse_diseno(BASE / "Preguntas - Diseño.docx", "Ingeniería de Software")
    todas += parse_bold_marks(
        BASE / "Preguntas - Parte 1.docx",
        "Ingeniería de Software",
        "Experiencia de Usuario (UX)",
        "Fundamentos de UX",
    )
    todas += parse_bold_marks(
        BASE / "Preguntas - Parte 2.docx",
        "Ingeniería de Software",
        "Experiencia de Usuario (UX)",
        "Heurísticas de Nielsen y prototipado",
    )

    out = BASE / "preguntas.json"
    out.write_text(json.dumps(todas, ensure_ascii=False, indent=2), encoding="utf-8")

    por_materia: dict[str, int] = {}
    for q in todas:
        por_materia[q["materia"]] = por_materia.get(q["materia"], 0) + 1
    print(f"Total preguntas: {len(todas)}")
    for m, n in por_materia.items():
        print(f"  {m}: {n}")
    print(f"Escrito en {out}")


if __name__ == "__main__":
    main()
